import json
from datetime import date, datetime, time
from pathlib import Path
from openpyxl import load_workbook

SOURCE = Path("Ekim 2026 Data.xlsx")
OUTPUT = Path("performance-data.json")

def clean(value):
    if value is None:
        return None
    if isinstance(value, (datetime, date, time)):
        return value.isoformat()
    if isinstance(value, (str, int, float, bool)):
        return value
    return str(value)

def main():
    if not SOURCE.exists():
        raise FileNotFoundError(f"Excel dosyası bulunamadı: {SOURCE}")

    # data_only=True reads Excel's last saved formula results without editing formulas.
    workbook = load_workbook(SOURCE, data_only=True, read_only=False)
    result = {
        "source": SOURCE.name,
        "sheets": []
    }

    for sheet in workbook.worksheets:
        rows = []
        for row_idx, row in enumerate(sheet.iter_rows(), start=1):
            cells = []
            for cell in row:
                value = clean(cell.value)
                if value is not None and value != "":
                    cells.append({
                        "column": cell.column,
                        "address": cell.coordinate,
                        "value": value,
                        "number_format": cell.number_format
                    })
            if cells:
                rows.append({"row": row_idx, "cells": cells})

        result["sheets"].append({
            "name": sheet.title,
            "max_row": sheet.max_row,
            "max_column": sheet.max_column,
            "rows": rows
        })

    OUTPUT.write_text(json.dumps(result, ensure_ascii=False, indent=2), encoding="utf-8")
    print(f"{OUTPUT} oluşturuldu. Sayfa sayısı: {len(result['sheets'])}")

if __name__ == "__main__":
    main()
